import re
import sys

XLSX_SKILL_DIR = r"C:/Users/Arinstne/.zcode/cli/plugins/cache/zcode-plugins-official/spreadsheets/0.1.8/skills/xlsx"
for sub in [XLSX_SKILL_DIR, XLSX_SKILL_DIR + "/templates"]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from base import (  # noqa: E402
    FORMATS,
    setup_sheet, style_header_row, style_data_row,
    auto_fit_columns, auto_fit_row_heights, font_caption,
)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment  # noqa: E402

BASE = r"E:/dev/NLP/Rewrite_terra/structure"
DAT = BASE + "/seen08000.dat.txt"
GEI = BASE + "/gameexe_chs.ini"
OUT = BASE + "/speaker_labels.xlsx"
OUT_FALLBACK = BASE + "/speaker_labels_filled.xlsx"
TSV_FULL = BASE + "/namae_map.tsv"
TSV_TERRA = BASE + "/speaker_map.tsv"

GAME_VERIFIED = {"瑚太朗＿必要": "瑚太朗", "男＿ルイス": "男人", "路易斯": "路易斯"}


def read_text(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8", "gbk", "cp932"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


# ---- 1. #NAMAE mapping table from Chinese Gameexe.ini ----
namae = {}       # label -> display
namae_line = {}  # label -> ini line no
namae_order = []
for i, ln in enumerate(read_text(GEI).splitlines(), 1):
    m = re.match(r'^#NAMAE\s*=\s*"([^"]*)"\s*,\s*"([^"]*)"\s*,\s*(.*)$', ln)
    if m:
        lb, disp = m.group(1), m.group(2)
        namae[lb] = disp
        namae_line[lb] = i
        namae_order.append((lb, disp, m.group(3).strip()))
print(f"#NAMAE entries: {len(namae)}")

with open(TSV_FULL, "w", encoding="utf-8", newline="") as f:
    f.write("# Rewrite 汉化版 gameexe.chs 的 #NAMAE 表（原始标签 -> 游戏内显示名），"
            "由 siglus-ssu -x --gei 还原\n")
    f.write("原始标签\t游戏显示名\t标志位\n")
    for lb, disp, flags in namae_order:
        f.write(f"{lb}\t{disp}\t{flags}\n")

# ---- 2. Terra scene speaker labels ----
with open(DAT, encoding="utf-8") as f:
    lines = f.read().splitlines()
str_defs = {}
for i in range(16, 20917):
    m = re.match(r"^\[(\d+)\] '(.*)'$", lines[i])
    if m:
        str_defs.setdefault(m.group(2), []).append((i + 1, int(m.group(1))))
occ = {}
for i, ln in enumerate(lines):
    if "CD_NAME " in ln:
        m = re.search(r'CD_NAME "(.*)"', ln)
        if m:
            occ.setdefault(m.group(1), []).append(i + 1)

rows = []
covered = 0
for lb, xs in occ.items():
    defs = str_defs.get(lb, [])
    if lb in namae:
        display = namae[lb]
        covered += 1
        if lb in GAME_VERIFIED and GAME_VERIFIED[lb] == display:
            note = "游戏确认 ✓ 且与 #NAMAE 表一致"
        else:
            note = f"来自 gameexe #NAMAE 表（ini 第 {namae_line[lb]} 行）"
    else:
        display = lb  # 无条目 → 引擎按原样显示
        note = "无 #NAMAE 条目，引擎按原样显示（建议抽查 1 次）"
    rows.append({
        "label": lb, "first": xs[0], "last": xs[-1], "uses": len(xs),
        "ndef": len(defs),
        "deffirst": defs[0][0] if defs else None,
        "deflast": defs[-1][0] if defs else None,
        "display": display, "note": note,
    })
rows.sort(key=lambda r: r["first"])
by_count = sorted(rows, key=lambda r: -r["uses"])
rank = {r["label"]: i + 1 for i, r in enumerate(by_count)}
print(f"Terra labels: {len(rows)}, covered by #NAMAE: {covered}")

with open(TSV_TERRA, "w", encoding="utf-8", newline="") as f:
    f.write("原始标签\t游戏显示名\t出现次数\t来源\n")
    for r in sorted(rows, key=lambda r: -r["uses"]):
        src = "NAMAE表" if r["note"].startswith(("来自", "游戏确认")) else "原样显示"
        f.write(f"{r['label']}\t{r['display']}\t{r['uses']}\t{src}\n")

# ---- 3. xlsx ----
wb = Workbook()
ws = wb.active
ws.title = "说话人标签"
headers = ["CD_NAME 首现行", "CD_NAME 末现行", "原始标签", "出现次数", "次数排名",
           "str 定义数", "str_list 首行", "str_list 末行", "游戏显示名", "备注"]
last_col = len(headers) + 1
setup_sheet(ws, title="Rewrite Terra 线说话人对照表（显示名已按 gameexe #NAMAE 表填充）", last_col=last_col)
for col_idx, h in enumerate(headers, start=2):
    ws.cell(row=4, column=col_idx, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=last_col)

NUM_COLS = {2, 3, 5, 6, 7, 8, 9}
for i, r in enumerate(rows):
    rn = 5 + i
    vals = [r["first"], r["last"], r["label"], r["uses"], rank[r["label"]],
            r["ndef"], r["deffirst"], r["deflast"], r["display"], r["note"]]
    for col_idx, v in enumerate(vals, start=2):
        ws.cell(row=rn, column=col_idx, value=v)
    style_data_row(ws, row_num=rn, col_start=2, col_end=last_col, row_index=i)
    for col_idx in NUM_COLS:
        c = ws.cell(row=rn, column=col_idx)
        c.number_format = FORMATS["integer"]
        c.alignment = Alignment(horizontal="right", vertical="center")

last_data_row = 4 + len(rows)
ws.auto_filter.ref = f"B4:K{last_data_row}"
ws.freeze_panes = "A5"
cap = ws.cell(row=last_data_row + 2, column=2,
              value=f"显示名来源：gameexe_chs.ini 的 #NAMAE 表（186 条，用 siglus-ssu -x --gei --angou 游戏目录 还原）；"
                    f"覆盖 {covered}/{len(rows)} 个 Terra 标签，其余无条目者引擎按原标签显示。"
                    "默认按 CD_NAME 首现行升序（游戏流程顺序），点表头箭头可改排序。"
                    "全量表见 namae_map.tsv，管线用映射见 speaker_map.tsv。")
cap.font = font_caption()
ws.merge_cells(start_row=last_data_row + 2, start_column=2, end_row=last_data_row + 2, end_column=last_col)
auto_fit_columns(ws, min_width=8, max_width=28, header_row=4, data_start_row=5)
auto_fit_row_heights(ws, header_row=4, data_start_row=5)
wb.properties.creator = "Z.ai"
try:
    wb.save(OUT)
    saved = OUT
except PermissionError:
    wb.save(OUT_FALLBACK)
    saved = OUT_FALLBACK
print(f"saved: {saved}")
